# -*- coding: utf-8 -*-
"""Build 简历-v7.docx.

内容基线 = 简历-v6.docx 的正文；本版新增/改动见 简历-v7-这版改了什么.md。
排版沿用 build_resume_docx.py 的同一套函数与数值（微软雅黑、A4、上下 1.1cm / 左右 1.4cm），
这样 v7 和之前的版本视觉上完全一致。

用法：
    D:/develop/python/python.exe build_resume_v7.py
"""

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DOCX = os.path.join(OUT_DIR, "简历-v7.docx")
DOC_TITLE = "吴枫 - Java 后端开发 - 个人简历 v7"

FONT = "微软雅黑"
BLACK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x1F, 0x38, 0x64)
STEEL = RGBColor(0x2E, 0x5C, 0x8A)
GRAY = RGBColor(0x59, 0x59, 0x59)

PAGE_W_CM = 21.0
PAGE_H_CM = 29.7
MARGIN_CM = 1.4
TOP_CM = 1.1
BOTTOM_CM = 1.1
USABLE_PT = (PAGE_W_CM - 2 * MARGIN_CM) / 2.54 * 72

NAME = "吴枫"
CONTACT_1 = "Java 后端开发 ｜ 电话 17338701320 ｜ 年龄 30"

OVERVIEW = (
    "Java 后端开发，主要做协同办公与系统集成交付，客户以国企和大型集团为主。熟悉"
    "<b>蓝凌 EKP 平台二次开发</b>与 <b>LBPM 流程引擎</b>事件扩展，完成过金蝶云星瀚、SAP、钉钉、"
    "IDaaS 等异构系统的接口对接；并用 <b>Spring Boot</b> 自建工程验证过跨系统接口的"
    "<b>幂等</b>与<b>并发聚合</b>方案。"
)

SKILLS = [
    ("Java 并发",
     "ThreadPoolExecutor 线程池（核心线程数、阻塞队列、拒绝策略）、Future 超时控制与单系统失败降级、"
     "ConcurrentHashMap；了解 CompletableFuture 并行编排与 CountDownLatch"),
    ("缓存与 Redis",
     "token 与字典缓存的周期复用与过期策略、SETNX 占位 + TTL 防止重复提交与重复推送、"
     "Lua 脚本保证比对与删除的原子性、空值缓存防穿透；了解穿透 / 击穿 / 雪崩的区别与常用对策"),
    ("数据库",
     "Oracle / SQL Server / MySQL；多表关联与 JOIN 去重、外键改软关联、覆盖索引与最左匹配、"
     "会用 EXPLAIN 看 type/key/rows/Extra、批量插入与更新、中间库多数据源配置"),
    ("平台与框架",
     "蓝凌 EKP 协同办公平台二次开发：表单建模、LBPM 流程引擎事件扩展"
     "（IEventListener / EndedTaskEvent）、门户组件 Portlet、权限模型；"
     "Spring Boot / Spring MVC / MyBatis；JSP 与 JSTL"),
    ("系统集成",
     "金蝶云星瀚（REST 接口、multipart/form-data 多附件上传、token 周期性缓存与鉴权）、"
     "SAP（XML 凭证生成、SFTP 上传、回执状态轮询）、钉钉开放平台（待办推送与审批状态回调）、"
     "IDaaS 单点登录（过滤器与 token 解析）、HTTP Basic 认证、EOP 供应商主数据同步"),
    ("工程与交付",
     "Linux、Git、Maven、SFTP、Docker 基础使用；版本发布与补丁包交付；"
     "需求说明书、开发说明文档、工作量评估与验收交付物输出"),
]

JOBS = [
    ("外包交付开发（承接蓝凌生态交付项目）", "2024.09 - 至今",
     "承接国企与大型集团客户的协同办公与系统集成交付，独立完成需求评估、开发、部署联调与交付文档输出。"
     "客户包括深圳巴士集团、深圳市融资租赁（集团）、罗格朗（惠州）、飞的数字化。"),
    ("深圳市蓝凌软件股份有限公司 · 交付部 ｜ Java 交付开发工程师", "2020.06 - 2024.09",
     "负责蓝凌 EKP 协同办公平台的国企与大型集团交付项目，覆盖模块开发、跨系统联调与上线支持。"
     "客户包括深圳能源集团、康佳集团、深圳巴士集团、深圳市融资租赁公司。"),
    ("中软国际（华为外包）｜ Java 开发工程师", "2019.07 - 2020.05",
     "中国移动音频彩铃 USDP 平台纯后端开发，使用 Oracle 存储过程完成十余个 Story。"),
    ("厦门同创空间信息技术有限公司 ｜ Java 开发工程师", "2018.03 - 2019.06",
     "开思时代小马物流平台，参与两个后端服务开发，负责甲方信息附件接口对接与联调。"),
]

PROJECTS_RECENT = [
    ("深圳巴士集团 ｜ OA 合同管理与金蝶云星瀚集成", "2024.09 - 2025.02", [
        "负责采购合同附件接收接口，基于 <b>REST 接口</b> + <b>multipart/form-data 多附件上传</b> + "
        "<b>JSON 表单解析</b>，实现采购侧合同与多附件落入 OA，解决采购合同附件无法进入 OA 的问题；"
        "单次接收上限 <b>{待补：附件数 / 总大小}</b>。",
        "负责合同字段与附件反向推送采购系统，基于 <b>token 周期性缓存 + 鉴权复用</b>，"
        "把凭证换取从每次调用一次降为 <b>{待补：换取频率}</b>，避免重复换取凭证的额外开销。",
        "负责合同台账导出改造，基于 <b>多表 JOIN + 主键去重</b>，消除左连接产生的重复行，"
        "导出重复数据 <b>{待补：重复行数} → 0</b>。",
        "负责合同权限申请公文流程与合同申请表单字段定制，基于 <b>流程事件推送待办</b>，"
        "打通合同权限申请与审批闭环。",
    ]),
    ("深圳市融资租赁（集团） ｜ 智慧办公系统二期", "2025.01 - 2025.09", [
        "开发会议通知监听器，基于 <b>LBPM 事件扩展（IEventListener + EndedTaskEvent）</b>，"
        "在会议节点结束时触发平台通知，实现会议通知自动下发。",
        "负责门户组件 <b>Portlet</b> 开发，实现工作安排与上会安排门户展示、领导批示定制与公文报表。",
        "负责议题管理模块，基于 <b>议题申请表字段与审批走向联动</b> + 上会安排 + 会议议题与公文纪要联动，"
        "解决议题、会议与公文数据割裂的问题。",
        "负责会议管理改造，实现 <b>会议室保管者字段由单值改为 List 多选</b>、会议信息传阅与会议室预约优化。",
    ]),
    ("深圳市融资租赁（集团） ｜ OA 待办推送钉钉工作台", "2025.12 - 2026.04", [
        "负责 OA 待办与钉钉工作台集成，基于 <b>钉钉开放平台接口 + 待办推送 + 状态回调</b>，"
        "实现待办推送、点击跳转 OA 原审批表单、审批完成后回调同步 OA，"
        "解决用户在钉钉内无法处理 OA 审批的问题。",
        "基于 <b>幂等键校验</b>避免同一待办重复推送，降低钉钉侧重复待办与用户投诉；"
        "日均推送 <b>{待补：条数}</b>，重复推送率 <b>{待补}%</b>。",
        "处理移动端与 PC 端跳转链接不一致的问题，统一 <b>{待补：链接拼接方案}</b>。",
    ]),
    ("罗格朗（惠州） ｜ OA 费控与 SAP 集成", "2025.02 - 2026.04", [
        "负责报销凭证模块，基于 <b>流程结束监听器 + XML 凭证生成 + SFTP 上传 + 回执状态轮询</b>，"
        "实现凭证归档与跨系统状态同步，解决 OA 与 SAP 状态不一致的问题；"
        "日均 <b>{待补：笔数}</b>，SFTP 上传成功率 <b>{待补}%</b>。",
        "负责报销凭证批量归档链路优化：定位到 <b>SFTP 工具类单例 + 方法级 synchronized</b> 导致批量上传全局串行、"
        "批量耗时随凭证数线性增长，改用 <b>ThreadPoolExecutor 并发上传 + 每线程独立 SFTP 连接隔离 + "
        "单张凭证失败隔离重试</b>，批量归档任务耗时 <b>{待补：串行 → 并发}</b>。",
        "负责成本中心模块，基于 <b>OA 组织架构与 SAP 成本中心对照映射</b> + 中间库（SQL Server）读取 + "
        "<b>大/小管理员二级权限模型</b>，实现成本中心公司关联与权限向下带出。",
        "负责费用事前申请单，基于 <b>逐级金额约束校验</b>（批准金额不高于申请金额、报销金额不高于批准金额），"
        "防止超申请审批与超批准报销。",
        "负责预算占用与释放协同，基于 <b>驳回/废弃状态分别触发释放调用 + 幂等键去重</b>，避免预算重复占用；"
        "幂等键 <b>{待补：键值 / 存储位置}</b>。",
    ]),
    ("其他交付项目 ｜ 合同费控与督办优化", "2025.11 - 2026.09", [
        "<b>飞的数字化 · 合同变更与费控</b>：实现合同变更带入<b>费控台账</b>"
        "（金额增加走从合同、金额减少走变更协议并同步更新合同金额）、"
        "合同付款权限带出与实际支付金额不高于未付金额校验，以及相对方变更与 EOP 供应商信息同步。",
        "<b>深圳巴士集团 · 督办模块优化</b>：引入 sys/task 产品包定制至 ekp/task，"
        "定位并修复平台二开引入流程后重复调用 save 引发的<b>数据库锁冲突</b>，改造 DAO 层 update 解决；"
        "同时修复<b>定时任务死循环</b>问题。",
    ]),
]

PROJECTS_EARLY = [
    ("深圳能源集团 ｜ 协同办公与合同管理系统", "2020.07 - 2021.12", [
        "负责合同管理模块交付开发，基于 <b>EKP 表单建模</b>与自定义流程实现合同起草、审批、用印与归档，"
        "并定制合同台账查询与导出，解决合同台账与审批数据分散的问题。",
        "负责合同台账大表查询优化，通过 <b>覆盖索引</b>与 JOIN 去重改造替换原多层嵌套查询，"
        "解决台账列表查询超时问题；<b>{待补：改造前耗时 → 改造后耗时、台账表数据量级}</b>。",
        "负责组织架构与权限模型定制，基于 EKP 组织架构与角色权限实现按部门、岗位的数据可见范围控制，"
        "解决跨部门合同数据越权可见的问题。",
    ]),
    ("康佳集团 ｜ 集团门户与多系统集成", "2021.12 - 2023.03", [
        "负责统一待办聚合门户组件开发，基于 <b>ThreadPoolExecutor 线程池</b>并发拉取多个业务系统待办并聚合展示，"
        "解决原串行拉取导致门户加载缓慢的问题；<b>{待补：串行耗时 → 并发后耗时、涉及系统数}</b>。",
        "负责跨系统数据同步接口开发，基于 token 周期性缓存与 HTTP Basic 认证完成与外部系统的数据交互，"
        "解决重复换取凭证带来的性能损耗。",
        "负责接口防重处理，按 <b>业务单据号</b>做幂等校验，避免同一单据重复提交与重复同步。",
    ]),
    ("深圳市融资租赁（集团） ｜ 智慧办公系统一期", "2023.03 - 2024.09", [
        "负责公文管理与会议管理模块交付开发，实现收发文表单定制、会议安排与会议室预约，"
        "解决公文与会议线下流转不可追溯的问题。",
        "负责流程引擎二次开发，基于 <b>LBPM 流程引擎</b>实现会签、加签、备选人与修改办理人能力，"
        "解决多级审批与办理人变更场景下的流程流转需求。",
        "负责批量数据导入导出开发，基于 Excel 模板与<b>线程池并发分片</b>处理批量数据，"
        "解决大批量数据单线程串行导入耗时过长的问题；<b>{待补：一万行耗时 before → after、线程数}</b>。",
    ]),
]

# ==================== 排版函数（与 build_resume_docx.py 一致） ====================

def set_run(run, size, bold=False, color=BLACK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def segments(text):
    """把 <b>..</b> 标记拆成 (文本, 是否加粗)。"""
    out = []
    pos = 0
    for m in re.finditer(r"<b>(.*?)</b>", text, re.S):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out


def add_para(doc, text="", size=9.2, bold=False, line=13.0, before=0, after=2,
             indent=0, hanging=0, right_tab=False, keep_next=False, bullet=False,
             color=BLACK, bullet_color=STEEL):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    pf.widow_control = True
    if keep_next:
        pf.keep_with_next = True
    if indent:
        pf.left_indent = Pt(indent)
    if hanging:
        pf.first_line_indent = Pt(-hanging)
    if right_tab:
        pf.tab_stops.add_tab_stop(Pt(USABLE_PT), WD_TAB_ALIGNMENT.RIGHT)
    if bullet:
        set_run(p.add_run("\u00b7 "), size, False, bullet_color)
    for chunk, is_bold in segments(text):
        if chunk:
            set_run(p.add_run(chunk), size, bold or is_bold, color)
    if not text:
        set_run(p.add_run(""), size, bold, color)
    return p


def add_header(doc, title, date, size=9.5, line=12.8):
    """标题靠左加粗、日期靠右灰色，中间用右对齐制表位。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(line)
    pf.keep_with_next = True
    pf.widow_control = True
    pf.tab_stops.add_tab_stop(Pt(USABLE_PT), WD_TAB_ALIGNMENT.RIGHT)
    set_run(p.add_run(title), size, True)
    set_run(p.add_run("\t"), size, True)
    set_run(p.add_run(date), size, False, GRAY)
    return p


def lighten_title_style(doc):
    style = doc.styles["Title"]
    style.font.color.rgb = BLACK
    pPr = style.element.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(PAGE_W_CM)
    section.page_height = Cm(PAGE_H_CM)
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)
    section.top_margin = Cm(TOP_CM)
    section.bottom_margin = Cm(BOTTOM_CM)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = BLACK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT)


def section_heading(doc, text, first=False):
    return add_para(doc, text, size=12.5, bold=True, line=15.5,
                    before=0 if first else 8.5, after=4.5, keep_next=True, color=NAVY)


def skills_block(doc):
    for direction, detail in SKILLS:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2.4)
        pf.line_spacing = Pt(13.0)
        pf.widow_control = True
        set_run(p.add_run(direction), 9.2, True, STEEL)
        set_run(p.add_run("：" + detail), 9.2, False, BLACK)


def group_heading(doc, text):
    return add_para(doc, text, size=10, bold=True, line=13, before=4.5, after=3, keep_next=True)


def bullets_block(doc, bullets):
    for b in bullets:
        add_para(doc, b, size=9.2, line=12.6, after=1.6, indent=11, hanging=11, bullet=True)
    add_para(doc, "", size=1, line=1, after=0)


# ==================== 生成 ====================

def build():
    doc = Document()
    setup_page(doc)
    lighten_title_style(doc)

    name = add_para(doc, NAME, size=20.5, bold=True, line=24, after=1.5, color=NAVY)
    name.style = doc.styles["Title"]
    for run in name.runs:
        set_run(run, 20.5, True, NAVY)

    add_para(doc, CONTACT_1, size=9.2, line=12.6, after=2, color=GRAY)

    section_heading(doc, "个人概述", first=True)
    add_para(doc, OVERVIEW, size=9.2, line=13.0, after=2)

    section_heading(doc, "专业技能")
    skills_block(doc)

    section_heading(doc, "工作经历")
    for title, date, desc in JOBS:
        add_header(doc, title, date)
        add_para(doc, desc, size=9, line=12.4, after=4)

    section_heading(doc, "项目经历")
    group_heading(doc, "2024.09 - 至今 ｜ 外包交付（承接蓝凌生态交付项目）")
    for title, date, bullets in PROJECTS_RECENT:
        add_header(doc, title, date)
        bullets_block(doc, bullets)

    group_heading(doc, "2020.06 - 2024.09 ｜ 蓝凌软件交付部")
    for title, date, bullets in PROJECTS_EARLY:
        add_header(doc, title, date)
        bullets_block(doc, bullets)

    section_heading(doc, "教育背景")
    add_para(doc, "2018.07 ｜ 泉州师范学院 ｜ 软件工程 ｜ 本科（2014.09 - 2018.07）", size=9.2, line=13.0)

    doc.core_properties.title = DOC_TITLE
    doc.core_properties.author = "吴枫"
    doc.save(OUT_DOCX)
    print("[output] " + OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    build()
