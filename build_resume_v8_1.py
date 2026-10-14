"""从 v8 定点修改并保留现有段落格式，不覆盖原文件。"""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parent

def build():
    doc = Document(ROOT / '简历-v8.docx')
    changes = {
        'Java 并发：': 'Java 并发：ThreadPoolExecutor 七参数（核心线程数、最大线程数、存活时间、时间单位、工作队列、线程工厂、拒绝策略）与执行流程；有界队列与拒绝处理；Future 超时与局部降级（自建工程验证）；synchronized 锁对象与粒度；ConcurrentHashMap、AtomicLong；了解 CompletableFuture、CountDownLatch 与 AQS',
        'Spring 体系：': 'Spring 体系：Spring Boot / Spring MVC / MyBatis（XML Mapper、#{} 与 ${}）；了解 @Transactional 事务边界、自调用、异常被吞及代理方式与方法可见性限制；Bean 生命周期与循环依赖；JSP 与 JSTL',
        '· 负责报销凭证批量归档链路优化：': '· 负责报销凭证批量归档链路分析：定位 SFTP 工具类单例与方法级 synchronized 的串行瓶颈；拟采用 ThreadPoolExecutor、独立连接与单张失败隔离，改造落地及耗时结果待核实。',
    }
    changed=[]
    for p in doc.paragraphs:
        for prefix, replacement in changes.items():
            if p.text.startswith(prefix):
                old=p.text
                # 沿用首个正文 run 的格式，保留段落与分页设置。
                runs=p.runs
                if runs:
                    runs[0].text=replacement
                    for run in runs[1:]: run.text=''
                else: p.add_run(replacement)
                changed.append((old,replacement))
                break
    assert len(changed)==3, changed
    doc.core_properties.title='吴枫 Java 后端开发 个人简历 v8.1'
    doc.save(ROOT / '简历-v8.1.docx')
    # 校验受保护章节逐字不变。
    source=Document(ROOT/'简历-v8.docx')
    def section(d,start,end=None):
        texts=[p.text for p in d.paragraphs]; a=texts.index(start)
        b=texts.index(end,a+1) if end else len(texts)
        return texts[a:b]
    for start,end in [('个人概述','专业技能'),('工作经历','项目经历'),('教育背景',None)]:
        assert section(source,start,end)==section(doc,start,end)
    print('生成 v8.1：3 处定点修改；受保护章节逐字一致')
if __name__=='__main__': build()
